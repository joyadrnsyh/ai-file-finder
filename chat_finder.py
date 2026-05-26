# UI modern + animasi sederhana untuk AI File Finder
# Install:
# pip install customtkinter sentence-transformers scikit-learn pypdf python-docx openpyxl pandas xlrd

import customtkinter as ctk
from pathlib import Path
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity
from pypdf import PdfReader
from docx import Document
import openpyxl
import pandas as pd
import numpy as np
import os
import shutil
import threading
import time

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

model = None
data_file = []
embeddings = None

FOLDER_TARGET = [
    Path.home() / "Documents",
    Path.home() / "Downloads",
    Path.home() / "Desktop"
]

SUPPORTED_READ_EXTENSIONS = [
    ".pdf", ".docx", ".txt", ".xlsx", ".xls", ".csv",
    ".py", ".html", ".css", ".js", ".php", ".java", ".cpp", ".json", ".xml"
]

CATEGORY_MAP = {
    "Dokumen": [".pdf", ".doc", ".docx", ".txt", ".ppt", ".pptx", ".xls", ".xlsx", ".csv"],
    "Gambar": [".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg"],
    "Video": [".mp4", ".mov", ".avi", ".mkv"],
    "Audio": [".mp3", ".wav", ".m4a"],
    "Archive": [".zip", ".rar", ".7z"],
    "Kode": [".py", ".html", ".css", ".js", ".php", ".java", ".cpp", ".json", ".xml"],
}


def baca_isi_file(file):
    ext = file.suffix.lower()

    try:
        if ext in [".txt", ".py", ".html", ".css", ".js", ".php", ".java", ".cpp", ".json", ".xml"]:
            return file.read_text(encoding="utf-8", errors="ignore")

        if ext == ".pdf":
            reader = PdfReader(str(file))
            text = ""
            for page in reader.pages[:8]:
                text += page.extract_text() or ""
            return text

        if ext == ".docx":
            doc = Document(str(file))
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        if ext == ".xlsx":
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
            text = ""
            for sheet in workbook.worksheets[:3]:
                text += f"\nSheet: {sheet.title}\n"
                for row in sheet.iter_rows(max_row=80, values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    if row_text.strip():
                        text += row_text + "\n"
            workbook.close()
            return text

        if ext == ".xls":
            excel_data = pd.read_excel(file, sheet_name=None, nrows=80)
            text = ""
            for sheet_name, df in list(excel_data.items())[:3]:
                text += f"\nSheet: {sheet_name}\n"
                text += df.astype(str).to_string(index=False)
            return text

        if ext == ".csv":
            df = pd.read_csv(file, nrows=100, encoding_errors="ignore")
            return df.astype(str).to_string(index=False)

    except Exception:
        return ""

    return ""


class AIFileFinder(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("AI File Finder Pro")
        self.geometry("1200x760")
        self.minsize(1000, 650)

        self.pending_moves = []
        self.is_indexing = False

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.create_sidebar()
        self.create_main_area()

        self.add_ai_message(
            "Halo. Saya bisa mencari file berdasarkan nama dan isi dokumen lokal: PDF, Word, Excel, TXT, CSV, dan file kode."
        )

    def create_sidebar(self):
        self.sidebar = ctk.CTkFrame(self, width=270, corner_radius=0, fg_color="#0f172a")
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_propagate(False)

        self.logo = ctk.CTkLabel(
            self.sidebar,
            text="AI Finder",
            font=("Segoe UI", 30, "bold"),
            text_color="#ffffff"
        )
        self.logo.pack(pady=(35, 5))

        self.subtitle = ctk.CTkLabel(
            self.sidebar,
            text="Local Semantic Search",
            font=("Segoe UI", 13),
            text_color="#94a3b8"
        )
        self.subtitle.pack(pady=(0, 25))

        self.status_card = ctk.CTkFrame(self.sidebar, fg_color="#111827", corner_radius=18)
        self.status_card.pack(padx=18, pady=(0, 18), fill="x")

        self.status_title = ctk.CTkLabel(
            self.status_card,
            text="SYSTEM STATUS",
            font=("Segoe UI", 11, "bold"),
            text_color="#38bdf8"
        )
        self.status_title.pack(anchor="w", padx=15, pady=(14, 4))

        self.status = ctk.CTkLabel(
            self.status_card,
            text="Belum siap",
            text_color="#cbd5e1",
            wraplength=210,
            justify="left"
        )
        self.status.pack(anchor="w", padx=15, pady=(0, 12))

        self.progress = ctk.CTkProgressBar(self.status_card)
        self.progress.pack(padx=15, pady=(0, 15), fill="x")
        self.progress.set(0)

        self.btn_index = self.sidebar_button("Index Ulang", self.start_indexing)
        self.btn_organize = self.sidebar_button("Preview Organize", self.preview_organize)
        self.btn_execute = self.sidebar_button("Jalankan Organize", self.execute_organize, "#16a34a", "#15803d")
        self.btn_clear = self.sidebar_button("Bersihkan Chat", self.clear_chat, "#334155", "#475569")

        self.info = ctk.CTkLabel(
            self.sidebar,
            text="Folder dipindai:\nDocuments\nDownloads\nDesktop",
            font=("Segoe UI", 12),
            text_color="#64748b",
            justify="left"
        )
        self.info.pack(side="bottom", anchor="w", padx=22, pady=25)

    def sidebar_button(self, text, command, color="#2563eb", hover="#1d4ed8"):
        btn = ctk.CTkButton(
            self.sidebar,
            text=text,
            height=44,
            corner_radius=14,
            fg_color=color,
            hover_color=hover,
            font=("Segoe UI", 13, "bold"),
            command=command
        )
        btn.pack(padx=18, pady=7, fill="x")
        return btn

    def create_main_area(self):
        self.main = ctk.CTkFrame(self, fg_color="#020617", corner_radius=0)
        self.main.grid(row=0, column=1, sticky="nsew")
        self.main.grid_rowconfigure(1, weight=1)
        self.main.grid_columnconfigure(0, weight=1)

        self.header = ctk.CTkFrame(self.main, fg_color="#020617", height=82)
        self.header.grid(row=0, column=0, sticky="ew", padx=25, pady=(20, 0))
        self.header.grid_columnconfigure(0, weight=1)

        self.header_title = ctk.CTkLabel(
            self.header,
            text="Chat Assistant",
            font=("Segoe UI", 26, "bold"),
            text_color="#f8fafc"
        )
        self.header_title.grid(row=0, column=0, sticky="w")

        self.header_desc = ctk.CTkLabel(
            self.header,
            text="Cari file dengan bahasa natural. Contoh: “cari dokumen yang membahas LSTM Tokopedia”.",
            font=("Segoe UI", 13),
            text_color="#94a3b8"
        )
        self.header_desc.grid(row=1, column=0, sticky="w", pady=(4, 0))

        self.chat_box = ctk.CTkScrollableFrame(
            self.main,
            fg_color="#020617",
            corner_radius=0
        )
        self.chat_box.grid(row=1, column=0, sticky="nsew", padx=25, pady=(10, 10))
        self.chat_box.grid_columnconfigure(0, weight=1)

        self.input_outer = ctk.CTkFrame(self.main, fg_color="#0f172a", corner_radius=22)
        self.input_outer.grid(row=2, column=0, sticky="ew", padx=25, pady=(0, 22))
        self.input_outer.grid_columnconfigure(0, weight=1)

        self.entry = ctk.CTkEntry(
            self.input_outer,
            placeholder_text="Ketik perintah...",
            height=54,
            border_width=0,
            fg_color="#111827",
            text_color="#f8fafc",
            font=("Segoe UI", 15)
        )
        self.entry.grid(row=0, column=0, sticky="ew", padx=(14, 10), pady=12)
        self.entry.bind("<Return>", lambda event: self.handle_command())

        self.btn_send = ctk.CTkButton(
            self.input_outer,
            text="Kirim",
            width=110,
            height=54,
            corner_radius=18,
            font=("Segoe UI", 14, "bold"),
            command=self.handle_command
        )
        self.btn_send.grid(row=0, column=1, padx=(0, 14), pady=12)

    def safe_ui(self, func):
        self.after(0, func)

    def animate_progress(self):
        if self.is_indexing:
            current = self.progress.get()
            next_value = current + 0.025
            if next_value > 0.95:
                next_value = 0.15
            self.progress.set(next_value)
            self.after(120, self.animate_progress)

    def add_typing(self):
        frame = ctk.CTkFrame(self.chat_box, fg_color="#111827", corner_radius=18)
        frame.grid(sticky="w", padx=8, pady=7)

        label = ctk.CTkLabel(
            frame,
            text="AI sedang memproses...",
            text_color="#94a3b8",
            font=("Segoe UI", 13),
            padx=14,
            pady=10
        )
        label.pack()

        return frame

    def add_user_message(self, text):
        bubble = ctk.CTkLabel(
            self.chat_box,
            text=text,
            fg_color="#2563eb",
            text_color="white",
            corner_radius=18,
            padx=16,
            pady=12,
            wraplength=650,
            justify="left",
            font=("Segoe UI", 14)
        )
        bubble.grid(sticky="e", padx=8, pady=7)

    def add_ai_message(self, text):
        bubble = ctk.CTkLabel(
            self.chat_box,
            text=text,
            fg_color="#111827",
            text_color="#f8fafc",
            corner_radius=18,
            padx=16,
            pady=12,
            wraplength=760,
            justify="left",
            font=("Segoe UI", 14)
        )
        bubble.grid(sticky="w", padx=8, pady=7)

    def add_file_result(self, title, score, path, preview):
        card = ctk.CTkFrame(self.chat_box, fg_color="#0f172a", corner_radius=20)
        card.grid(sticky="ew", padx=8, pady=8)
        card.grid_columnconfigure(0, weight=1)

        top = ctk.CTkFrame(card, fg_color="transparent")
        top.grid(row=0, column=0, sticky="ew", padx=16, pady=(14, 5))
        top.grid_columnconfigure(0, weight=1)

        name = ctk.CTkLabel(
            top,
            text=title,
            font=("Segoe UI", 15, "bold"),
            text_color="#f8fafc",
            anchor="w"
        )
        name.grid(row=0, column=0, sticky="w")

        badge = ctk.CTkLabel(
            top,
            text=f"{score:.2f}",
            fg_color="#1d4ed8",
            text_color="white",
            corner_radius=12,
            padx=10,
            pady=4,
            font=("Segoe UI", 12, "bold")
        )
        badge.grid(row=0, column=1, sticky="e")

        location = ctk.CTkLabel(
            card,
            text=f"Lokasi: {path}",
            text_color="#94a3b8",
            font=("Segoe UI", 12),
            anchor="w",
            wraplength=760,
            justify="left"
        )
        location.grid(row=1, column=0, sticky="w", padx=16, pady=(0, 8))

        body = ctk.CTkLabel(
            card,
            text=f"Preview: {preview}",
            text_color="#cbd5e1",
            font=("Segoe UI", 13),
            anchor="w",
            wraplength=760,
            justify="left"
        )
        body.grid(row=2, column=0, sticky="w", padx=16, pady=(0, 12))

        actions = ctk.CTkFrame(card, fg_color="transparent")
        actions.grid(row=3, column=0, sticky="e", padx=16, pady=(0, 14))

        open_btn = ctk.CTkButton(
            actions,
            text="Buka File",
            width=100,
            height=34,
            corner_radius=12,
            command=lambda: os.startfile(path)
        )
        open_btn.pack(side="left", padx=5)

        folder_btn = ctk.CTkButton(
            actions,
            text="Buka Folder",
            width=110,
            height=34,
            corner_radius=12,
            fg_color="#334155",
            hover_color="#475569",
            command=lambda: os.startfile(path.parent)
        )
        folder_btn.pack(side="left", padx=5)

    def clear_chat(self):
        for widget in self.chat_box.winfo_children():
            widget.destroy()

    def handle_command(self):
        query = self.entry.get().strip()
        if not query:
            return

        self.entry.delete(0, "end")
        self.add_user_message(query)

        lower_query = query.lower()

        if "organize" in lower_query or "rapikan" in lower_query or "atur file" in lower_query:
            self.preview_organize()
            return

        if embeddings is None:
            self.add_ai_message("Index belum tersedia. Saya buat index dulu, lalu mencari file.")
            thread = threading.Thread(target=lambda: self.index_then_search(query), daemon=True)
            thread.start()
        else:
            thread = threading.Thread(target=lambda: self.search_file(query), daemon=True)
            thread.start()

    def start_indexing(self):
        if self.is_indexing:
            self.add_ai_message("Indexing masih berjalan.")
            return

        thread = threading.Thread(target=self.index_files, daemon=True)
        thread.start()

    def index_then_search(self, query):
        self.index_files()
        self.search_file(query)

    def index_files(self):
        global model, data_file, embeddings

        self.is_indexing = True
        self.safe_ui(lambda: self.btn_index.configure(state="disabled"))
        self.safe_ui(lambda: self.status.configure(text="Memuat model Transformer..."))
        self.safe_ui(lambda: self.progress.set(0.05))
        self.safe_ui(self.animate_progress)

        if model is None:
            model = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")

        temp_data = []
        total_scanned = 0
        total_readable = 0

        self.safe_ui(lambda: self.status.configure(text="Membaca file lokal..."))

        for folder in FOLDER_TARGET:
            if folder.exists():
                for file in folder.rglob("*"):
                    if file.is_file():
                        total_scanned += 1
                        isi_file = ""

                        if file.suffix.lower() in SUPPORTED_READ_EXTENSIONS:
                            isi_file = baca_isi_file(file)
                            if isi_file.strip():
                                total_readable += 1

                        teks_file = f"""
Nama file: {file.name}
Ekstensi: {file.suffix}
Folder: {file.parent.name}
Path: {file}
Isi file:
{isi_file[:4000]}
"""
                        temp_data.append({
                            "path": file,
                            "teks": teks_file,
                            "isi": isi_file[:1200]
                        })

        if not temp_data:
            self.safe_ui(lambda: self.add_ai_message("Tidak ada file ditemukan."))
            self.safe_ui(lambda: self.status.configure(text="Gagal: tidak ada file"))
            self.safe_ui(lambda: self.btn_index.configure(state="normal"))
            self.is_indexing = False
            return

        self.safe_ui(lambda: self.status.configure(text="Membuat embedding..."))

        teks = [item["teks"] for item in temp_data]
        temp_embeddings = model.encode(teks, show_progress_bar=False)

        data_file = temp_data
        embeddings = temp_embeddings

        self.is_indexing = False
        self.safe_ui(lambda: self.progress.set(1))
        self.safe_ui(lambda: self.btn_index.configure(state="normal"))
        self.safe_ui(lambda: self.status.configure(
            text=f"Siap\nFile: {len(data_file)}\nIsi terbaca: {total_readable}"
        ))
        self.safe_ui(lambda: self.add_ai_message(
            f"Index selesai.\nTotal file discan: {total_scanned}\nFile yang isinya terbaca: {total_readable}"
        ))

    def search_file(self, query):
        global model, data_file, embeddings

        if embeddings is None:
            self.safe_ui(lambda: self.add_ai_message("Index belum tersedia."))
            return

        typing = None
        self.safe_ui(lambda: None)

        query_embedding = model.encode([query])
        skor = cosine_similarity(query_embedding, embeddings)[0]
        ranking = np.argsort(skor)[::-1][:7]

        self.safe_ui(lambda: self.add_ai_message("Saya menemukan file yang paling relevan:"))

        for index in ranking:
            path = data_file[index]["path"]
            nilai = skor[index]
            preview = data_file[index].get("isi", "").strip()

            if preview:
                preview = preview[:520].replace("\n", " ")
            else:
                preview = "Tidak ada preview isi. File mungkin tidak didukung, kosong, atau hanya terbaca dari nama file."

            self.safe_ui(
                lambda path=path, nilai=nilai, preview=preview:
                self.add_file_result(path.name, nilai, path, preview)
            )

    def get_category(self, file):
        ext = file.suffix.lower()
        for category, extensions in CATEGORY_MAP.items():
            if ext in extensions:
                return category
        return "Lainnya"

    def preview_organize(self):
        self.pending_moves = []
        organize_root = Path.home() / "Documents" / "AI_Organized_Files"

        for folder in FOLDER_TARGET:
            if folder.exists():
                for file in folder.rglob("*"):
                    if file.is_file():
                        if "AI_Organized_Files" in str(file):
                            continue

                        category = self.get_category(file)
                        target_folder = organize_root / category
                        target_path = target_folder / file.name

                        if file.resolve() != target_path.resolve():
                            self.pending_moves.append((file, target_path, category))

        if not self.pending_moves:
            self.add_ai_message("Tidak ada file yang perlu dirapikan.")
            return

        self.add_ai_message(
            f"Saya menemukan {len(self.pending_moves)} file yang bisa dirapikan.\n"
            f"Folder tujuan: {organize_root}\n"
            f"Berikut preview 10 file pertama:"
        )

        for source, target, category in self.pending_moves[:10]:
            self.add_file_result(
                title=f"{source.name} → {category}",
                score=1.00,
                path=source,
                preview=f"Akan dipindahkan dari {source.parent} ke {target.parent}"
            )

        self.add_ai_message("Klik tombol 'Jalankan Organize' jika sudah yakin.")

    def execute_organize(self):
        if not self.pending_moves:
            self.add_ai_message("Belum ada preview organize. Klik 'Preview Organize' dulu.")
            return

        moved = 0
        skipped = 0

        for source, target, category in self.pending_moves:
            try:
                target.parent.mkdir(parents=True, exist_ok=True)

                final_target = target
                counter = 1

                while final_target.exists():
                    final_target = target.with_name(f"{target.stem}_{counter}{target.suffix}")
                    counter += 1

                shutil.move(str(source), str(final_target))
                moved += 1

            except Exception:
                skipped += 1

        self.pending_moves = []
        self.add_ai_message(
            f"Organize selesai.\nFile dipindahkan: {moved}\nFile dilewati: {skipped}"
        )

        self.start_indexing()


if __name__ == "__main__":
    app = AIFileFinder()
    app.mainloop()